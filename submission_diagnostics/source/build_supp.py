import copy, re, sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from refs import REFS, OLD
import content as C

# --- reproduce manuscript citation numbering ---
CITE_RE = re.compile(r'\[@([^\]]+)\]'); order = []
def collect(text):
    for m in CITE_RE.finditer(text):
        for k in m.group(1).split(';'):
            if k not in order: order.append(k.strip())
TABLES = {'T1': C.T1, 'T2': C.T2, 'T3': C.T3, 'T4': C.T4, 'T5': C.T5}
for kind, payload in C.BODY:
    if kind == 'p': collect(payload)
    elif kind == 'table':
        for r in TABLES[payload]['rows']:
            if payload == 'T2': collect('[@%s]' % r[0])
            else:
                for c in r: collect(c)
NUM = {k: i + 1 for i, k in enumerate(order)}

def remap(text):
    def f(m):
        nums = []
        for part in m.group(1).split(','):
            part = part.strip()
            if '–' in part or '-' in part:
                a, b = re.split('[–-]', part); rng = range(int(a), int(b) + 1)
            else: rng = [int(part)]
            nums += [NUM[OLD[n]] for n in rng]
        nums = sorted(set(nums))
        return '[' + ','.join(map(str, nums)) + ']'
    return re.sub(r'\[(\d+(?:\s*[,–-]\s*\d+)*)\]', f, text)

def clean(text):
    text = remap(text)
    text = text.replace('source: limited primary text', 'source: abstract and accessible primary material')
    text = text.replace('Source: indexed abstract/limited primary text', 'Source: abstract and accessible primary material')
    text = text.replace('source: complete primary report', 'source: full text').replace('Source: complete primary report', 'Source: full text')
    return text

src_ms = docx.Document('../orig/manuscript.docx')
src = docx.Document('../orig/supp_tables.docx')
def rows(t):
    out = []
    for r in t.rows:
        cells = []; prev = None
        for c in r.cells:
            if c._tc is prev: continue
            prev = c._tc; cells.append(c.text.strip())
        out.append(cells)
    return out
S1 = rows(src.tables[0]); S2A = rows(src.tables[1]); S2B = rows(src.tables[2]); S3 = rows(src.tables[3]); S4A = rows(src.tables[4]); S4B = rows(src.tables[5])
EXTR = rows(src_ms.tables[2])

# --- rebuild on the supplementary template ---
doc = docx.Document('../orig/supp_tables.docx')
body = doc.element.body; kids = list(body.iterchildren()); final = kids[-1]
for k in kids[:-1]: body.remove(k)
TOK = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*)')
def add_inline(p, text, size=None, bold=False):
    for part in TOK.split(text):
        if not part: continue
        b = bold; i = False
        if part.startswith('**'): part = part[2:-2]; b = True
        elif part.startswith('*'): part = part[1:-1]; i = True
        r = p.add_run(part); r.bold = b or None; r.italic = i or None
        if size: r.font.size = Pt(size)
def para(style, text='', **kw):
    p = doc.add_paragraph(style=style)
    if text: add_inline(p, text, **kw)
    return p

def table(caption, data, widths, size=7.5, note=None, bold_header=True):
    para('MDPI_2.2_heading2', caption)
    tb = doc.add_table(rows=len(data), cols=len(data[0]))
    tb.style = doc.styles['MDPI_4.1_three_line_table']
    tblPr = tb._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(w))
    for ri, row in enumerate(data):
        tr = tb.rows[ri]
        if ri == 0:
            th = OxmlElement('w:tblHeader'); tr._tr.get_or_add_trPr().append(th)
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            tcPr = cell._tc.get_or_add_tcPr(); tcW = OxmlElement('w:tcW'); tcW.set(qn('w:w'), str(widths[ci])); tcW.set(qn('w:type'), 'dxa'); tcPr.append(tcW)
            p = cell.paragraphs[0]; p.style = doc.styles['MDPI_4.2_table_body']; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, clean(val), size=size, bold=(ri == 0 and bold_header))
    if note: para('MDPI_3.2_text_no_indent', note, size=8.5)

para('MDPI_1.2_title', 'Supplementary Materials')
para('MDPI_3.2_text_no_indent', '**' + C.TITLE + '**')
para('MDPI_3.2_text_no_indent', 'Majed S. Alsanea, Yosef Hasan Jbara, Alshimaa Hamdy Ismail, Farha Mujeeb Ahmed Shaikh, and Mohammed Gh. Alzahrani')
para('MDPI_3.2_text_no_indent', 'Contents: Table S1. Information sources, search strategies, limits, dates, and yields. Table S2. Detailed data extraction for the 47 included reports. Table S3. Excluded reports that might appear eligible, with reasons. Table S4. Cohort-overlap map. Table S5. PROBAST+AI risk-of-bias judgments by domain. Table S6. PROBAST+AI applicability judgments. Reference numbers in square brackets correspond to the reference list of the main article.', size=9)

# S1
table('Table S1. Information sources, search strategies, limits, dates, and yields', S1, [2300, 9800, 1400, 1400],
      note='The nine bibliographic databases contributed 14,388 records before deduplication. Google Scholar, registries, and ProQuest brought the database/register route to 14,667 records; citation searching, preprint servers, hand-searching, and author contact contributed 245 records through other methods (total 14,912). After removal of 5433 duplicates, 9479 records were screened, 352 reports were sought, 11 were not retrieved, 341 were assessed in full, 294 were excluded, and 47 were included.')
# S2 detailed extraction
hdr = ['Report [ref.]', 'Population, cohort, sample, and fallers', 'Gait technology, assessment, and candidate predictors', 'Fall outcome, horizon, and ascertainment', 'Model and validation stage', 'Performance and calibration', 'PROBAST+AI overall judgment', 'Principal limitation and source']
ext = [hdr] + [[c.replace('†', ' †') for c in r] for r in EXTR[1:]]
table('Table S2. Detailed data extraction for the 47 included reports', ext, [1400, 2200, 2200, 1700, 1900, 2500, 1000, 2200], size=7,
      note='† Full text not accessible; data were extracted from the published abstract and accessible primary material, and unreported items were recorded as NR. Abbreviations are defined in the footnote to Table 2 of the main article. The PROBAST+AI column gives the overall development and evaluation judgments for the principal gait-containing model of each report (Tables S5 and S6).')
# S3 exclusions
para('MDPI_2.2_heading2', 'Table S3. Excluded reports that might appear eligible, with reasons')
para('MDPI_3.2_text_no_indent', 'Principal reasons for the 294 reports excluded after full-text assessment: no prospective fall outcome or no valid temporal separation, 99; fall detection or event recognition, 53; no individualized prediction-model performance, 44; no sensor-derived gait predictor, 35; age criterion not met, 27; ineligible publication type, 20; simulated or synthetic fall data, 10; duplicate report, 6. The reports below are those that might reasonably appear to meet the eligibility criteria.', size=9)
tb = doc.paragraphs[-1]
S2Bc = [S2B[0]] + S2B[1:]
table('', S2Bc, [6800, 2300, 5800], size=7.5)
doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p) if False else None
# S4 overlap
table('Table S4. Cohort-overlap map', S3, [2600, 4200, 5000, 3100], size=8,
      note='Seven multi-report cohort families covered 17 reports, and 30 single-report families contributed one report each. Because participant independence could not be established for Drover 2017, it was grouped with the Howcroft programme in the primary analysis (37 families); treating it as independent (38 families) did not change any conclusion.')
# S5 RoB
S5 = [['Report [ref.]', 'DOI', 'Dev P/D', 'Dev Pred', 'Dev Out', 'Dev Anal', 'Dev Overall', 'Eval P/D', 'Eval Pred', 'Eval Out', 'Eval Anal', 'Eval Overall', 'Rationale']] + S4A[1:]
table('Table S5. PROBAST+AI risk-of-bias judgments by domain', S5, [1500, 1900, 560, 560, 560, 560, 620, 560, 560, 560, 560, 620, 5000], size=7,
      note='Part A (Dev) assesses concern about model-development quality and Part B (Eval) risk of bias in model evaluation. P/D: participants and data sources; Pred: predictors; Out: outcome; Anal: analysis; H: high; L: low; U: unclear; NA: not applicable (no model developed). An overall low judgment required all domains to be low. Judgments refer to the principal gait-containing model or shared development/evaluation pipeline of each report.')
S6 = [['Report [ref.]', 'DOI', 'Dev P/D', 'Dev Pred', 'Dev Out', 'Dev Overall', 'Eval P/D', 'Eval Pred', 'Eval Out', 'Eval Overall', 'Rationale']] + S4B[1:]
table('Table S6. PROBAST+AI applicability judgments', S6, [1500, 1900, 600, 600, 600, 660, 600, 600, 600, 660, 5900], size=7,
      note='Applicability was judged against the review question: prediction of subsequent falls in a broadly defined older-adult population using wearable- or vision-derived gait predictors. Considerations included age, sex, ethnicity where reported, disease, care setting, and reproducibility of the complete sensor-to-model pipeline. Abbreviations as in Table S5.')

# drop empty heading created for S3 sub-table
for p in list(doc.paragraphs):
    if p.style.name == 'MDPI_2.2_heading2' and not p.text.strip():
        p._p.getparent().remove(p._p)
doc.core_properties.title = 'Supplementary Materials: ' + C.TITLE
doc.core_properties.author = 'Majed S. Alsanea; Yosef Hasan Jbara; Alshimaa Hamdy Ismail; Farha Mujeeb Ahmed Shaikh; Mohammed Gh. Alzahrani'
doc.core_properties.comments = ''; doc.core_properties.last_modified_by = 'Yosef Hasan Jbara'
os.makedirs('out', exist_ok=True)
doc.save('supp_tmp.docx')
from postproc import to_diagnostics
to_diagnostics('supp_tmp.docx','out/Supplementary_Materials.docx'); os.remove('supp_tmp.docx'); print('ok')
