import copy, re, sys, zipfile, shutil, os
sys.path.insert(0, os.path.dirname(__file__))
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from refs import REFS, INCLUDED
import content as C

SRC = '../orig/manuscript.docx'
OUT_TMP = 'ms_tmp.docx'
OUT = 'out/Manuscript_Diagnostics.docx'

# ---------------- citation numbering ----------------
CITE_RE = re.compile(r'\[@([^\]]+)\]')
order = []
def collect(text):
    for m in CITE_RE.finditer(text):
        for k in m.group(1).split(';'):
            k = k.strip()
            assert k in REFS, k
            if k not in order: order.append(k)
TABLES = {'T1': C.T1, 'T2': C.T2, 'T3': C.T3, 'T4': C.T4, 'T5': C.T5}
for kind, payload in C.BODY:
    if kind == 'p': collect(payload)
    elif kind == 'table':
        t = TABLES[payload]
        if payload == 'T2':
            for r in t['rows']: collect('[@%s]' % r[0])
        else:
            for r in t['rows']:
                for c in r: collect(c)
for k in REFS:
    assert k in order, 'uncited reference: ' + k
NUM = {k: i + 1 for i, k in enumerate(order)}

def fmt_nums(nums):
    nums = sorted(set(nums)); out = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        if j - i >= 2: out.append(f"{nums[i]}–{nums[j]}")
        else: out.extend(str(n) for n in nums[i:j + 1])
        i = j + 1
    return '[' + ','.join(out) + ']'

def resolve(text):
    return CITE_RE.sub(lambda m: fmt_nums([NUM[k.strip()] for k in m.group(1).split(';')]), text)

TOK = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|\^[^^]+\^)')
def add_inline(par, text, size=None, bold_all=False):
    text = resolve(text)
    for part in TOK.split(text):
        if not part: continue
        b = i = sup = False
        if part.startswith('**'): part = part[2:-2]; b = True
        elif part.startswith('*'): part = part[1:-1]; i = True
        elif part.startswith('^'): part = part[1:-1]; sup = True
        r = par.add_run(part)
        if b or bold_all: r.bold = True
        if i: r.italic = True
        if sup: r.font.superscript = True
        if size: r.font.size = Pt(size)
    return par

# ---------------- open template & harvest front matter ----------------
doc = docx.Document(SRC)
body = doc.element.body
kids = list(body.iterchildren())
final_sect = kids[-1]
assert final_sect.tag == qn('w:sectPr')
sect_paras = [k for k in kids if k.tag == qn('w:p') and k.find(qn('w:pPr')) is not None and k.find(qn('w:pPr')).find(qn('w:sectPr')) is not None]
assert len(sect_paras) == 2
sect_portrait = copy.deepcopy(sect_paras[0].find(qn('w:pPr')).find(qn('w:sectPr')))
sect_landscape = copy.deepcopy(sect_paras[1].find(qn('w:pPr')).find(qn('w:sectPr')))
def style_of(el):
    ps = el.find(qn('w:pPr'))
    if ps is None or ps.find(qn('w:pStyle')) is None: return None
    return ps.find(qn('w:pStyle')).get(qn('w:val'))
front = {}
for k in kids[:15]:
    st = style_of(k) if k.tag == qn('w:p') else 'TABLE'
    front.setdefault(st, []).append(copy.deepcopy(k))
for k in kids[:-1]: body.remove(k)

def append(el): final_sect.addprevious(el)
append(front['MDPI11articletype'][0])
tp = front['MDPI12title'][0]
for r in tp.findall(qn('w:r')): tp.remove(r)
append(tp)
p = docx.text.paragraph.Paragraph(tp, doc); p.add_run(C.TITLE)
append(front['MDPI13authornames'][0])
# front-matter floating table (academic editor / dates / copyright)
ft = front['TABLE'][0]
append(ft)
for a in front['MDPI16affiliation']: append(a)

def para(style, text='', **kw):
    p = doc.add_paragraph(style=style)
    if text: add_inline(p, text, **kw)
    return p

# abstract + keywords
p = para('MDPI_1.7_abstract'); r = p.add_run('Abstract:'); r.bold = True
for head, txt in C.ABSTRACT:
    p.add_run(' '); rr = p.add_run(head); rr.bold = True; add_inline(p, txt)
p = para('MDPI_1.8_keywords'); r = p.add_run('Keywords:'); r.bold = True; p.add_run(' ' + C.KEYWORDS)
doc.add_paragraph(style='MDPI_1.9_line')

# ---------------- tables & figures ----------------
def set_cell_width(cell, w):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None: tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(w)); tcW.set(qn('w:type'), 'dxa')

def add_table(t, size=9, align_right=True, left_cols=(0,), small=False):
    para('MDPI_4.1_table_caption', t['caption'])
    ncols = len(t['header'])
    tb = doc.add_table(rows=1 + len(t['rows']), cols=ncols)
    tb.style = doc.styles['MDPI_4.1_three_line_table']
    tblPr = tb._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'right' if align_right else 'center'); tblPr.append(jc)
    grid = tb._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), t['widths']): gc.set(qn('w:w'), str(w))
    rows = [t['header']] + t['rows']
    for ri, row in enumerate(rows):
        tr = tb.rows[ri]
        if ri == 0:
            trPr = tr._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
        for ci, val in enumerate(row):
            cell = tr.cells[ci]; set_cell_width(cell, t['widths'][ci])
            cp = cell.paragraphs[0]; cp.style = doc.styles['MDPI_4.2_table_body']
            if ci in left_cols or small: cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(cp, val, size=size, bold_all=(ri == 0))
    if t.get('footer'): para('MDPI_4.3_table_footer', t['footer'])

def t2_rows():
    rows = []
    for r in C.T2['rows']:
        key, dag = r[0], r[1]
        name = REFS[key].split(',')[0]
        year = re.search(r'\*\*(\d{4})\*\*', REFS[key]).group(1)
        study = f"{name} et al. {year} [@{key}]{dag}"
        rows.append([study] + r[2:])
    return rows

def add_figure(f):
    p = doc.add_paragraph(style='MDPI_5.2_figure'); p.add_run().add_picture(f['file'], width=Cm(f['width_cm']))
    cp = para('MDPI_5.1_figure_caption', f['caption'])
    if f['note']: add_inline(cp, ' ' + f['note'])

def sect_break(sectPr):
    p = doc.add_paragraph(); pPr = p._p.get_or_add_pPr(); pPr.append(copy.deepcopy(sectPr))

for kind, payload in C.BODY:
    if kind == 'h1': para('MDPI_2.1_heading1', payload)
    elif kind == 'h2': para('MDPI_2.2_heading2', payload)
    elif kind == 'p': para('MDPI_3.1_text', payload)
    elif kind == 'figure': add_figure(C.FIGS[payload])
    elif kind == 'table':
        if payload == 'T2':
            sect_break(sect_portrait)
            t = dict(C.T2); t['rows'] = t2_rows()
            add_table(t, size=7.5, align_right=False, small=True)
            sect_break(sect_landscape)
        else:
            add_table(TABLES[payload], size=9, left_cols=(0,) if payload != 'T5' else (0, 1))

# back matter
for head, txt in C.BACK:
    p = para('MDPI_6.2_back_matter'); r = p.add_run(head); r.bold = True; add_inline(p, txt)
para('MDPI_2.1_heading1', 'Abbreviations')
para('MDPI_3.2_text_no_indent', 'The following abbreviations are used in this manuscript: ' + '; '.join(f'{a}, {b}' for a, b in C.ABBREVIATIONS) + '.')
para('MDPI_2.1_heading1', 'References')
for k in order:
    para('MDPI_8.1_references', REFS[k])
doc.core_properties.title = C.TITLE
doc.core_properties.author = 'Majed S. Alsanea; Yosef Hasan Jbara; Alshimaa Hamdy Ismail; Farha Mujeeb Ahmed Shaikh; Mohammed Gh. Alzahrani'
doc.core_properties.comments = ''; doc.core_properties.keywords = C.KEYWORDS
doc.core_properties.last_modified_by = 'Yosef Hasan Jbara'
xml = doc.element.xml
for rid, rel in list(doc.part.rels.items()):
    if rel.reltype.endswith('/image') and ('r:embed="%s"' % rid) not in xml:
        doc.part.drop_rel(rid) if hasattr(doc.part,'drop_rel') else doc.part.rels.pop(rid)
doc.save(OUT_TMP)

from postproc import to_diagnostics
os.makedirs('out', exist_ok=True)
to_diagnostics(OUT_TMP, OUT); os.remove(OUT_TMP)
print('refs:', len(order)); print('saved', OUT)
