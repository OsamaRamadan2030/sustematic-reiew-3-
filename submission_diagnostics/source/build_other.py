import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import content as C

AUTH = 'Majed S. Alsanea; Yosef Hasan Jbara; Alshimaa Hamdy Ismail; Farha Mujeeb Ahmed Shaikh; Mohammed Gh. Alzahrani'

def base_doc():
    d = docx.Document()
    st = d.styles['Normal']; st.font.name = 'Palatino Linotype'; st.font.size = Pt(10.5)
    d.styles['Normal'].paragraph_format.space_after = Pt(4)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Palatino Linotype')
    for s in d.sections:
        s.page_width = Cm(21.0); s.page_height = Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2.0)
    d.core_properties.author = AUTH; d.core_properties.last_modified_by = 'Yosef Hasan Jbara'; d.core_properties.comments = ''
    import datetime; now=datetime.datetime(2026,9,29,12,0,0); d.core_properties.created=now; d.core_properties.modified=now; d.core_properties.revision=1
    return d

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcol); tcPr.append(sh)

# ---------------- PRISMA 2020 checklist ----------------
P = [
('TITLE', None, None),
('Title', '1', 'Identify the report as a systematic review.', 'Title'),
('ABSTRACT', None, None),
('Abstract', '2', 'See the PRISMA 2020 for Abstracts checklist.', 'Abstract (structured: background/objectives, methods, results, conclusions, registration)'),
('INTRODUCTION', None, None),
('Rationale', '3', 'Describe the rationale for the review in the context of existing knowledge.', 'Section 1, paragraphs 1–4'),
('Objectives', '4', 'Provide an explicit statement of the objective(s) or question(s) the review addresses.', 'Section 1, final paragraph; Section 2.1'),
('METHODS', None, None),
('Eligibility criteria', '5', 'Specify the inclusion and exclusion criteria for the review and how studies were grouped for the syntheses.', 'Section 2.2; Table 1; Section 2.8'),
('Information sources', '6', 'Specify all databases, registers, websites, organisations, reference lists and other sources searched or consulted to identify studies. Specify the date when each source was last searched or consulted.', 'Section 2.3; Table S1'),
('Search strategy', '7', 'Present the full search strategies for all databases, registers and websites, including any filters and limits used.', 'Table S1'),
('Selection process', '8', 'Specify the methods used to decide whether a study met the inclusion criteria of the review, including how many reviewers screened each record and each report retrieved, whether they worked independently, and if applicable, details of automation tools used in the process.', 'Section 2.4'),
('Data collection process', '9', 'Specify the methods used to collect data from reports, including how many reviewers collected data from each report, whether they worked independently, any processes for obtaining or confirming data from study investigators, and if applicable, details of automation tools used in the process.', 'Section 2.5'),
('Data items', '10a', 'List and define all outcomes for which data were sought. Specify whether all results that were compatible with each outcome domain in each study were sought, and if not, the methods used to decide which results to collect.', 'Sections 2.5, 2.6, and 2.8'),
('', '10b', 'List and define all other variables for which data were sought (e.g., participant and intervention characteristics, funding sources). Describe any assumptions made about any missing or unclear information.', 'Section 2.5'),
('Study risk of bias assessment', '11', 'Specify the methods used to assess risk of bias in the included studies, including details of the tool(s) used, how many reviewers assessed each study and whether they worked independently, and if applicable, details of automation tools used in the process.', 'Section 2.7'),
('Effect measures', '12', 'Specify for each outcome the effect measure(s) (e.g., risk ratio, mean difference) used in the synthesis or presentation of results.', 'Section 2.6'),
('Synthesis methods', '13a', 'Describe the processes used to decide which studies were eligible for each synthesis.', 'Section 2.8'),
('', '13b', 'Describe any methods required to prepare the data for presentation or synthesis, such as handling of missing summary statistics, or data conversions.', 'Sections 2.5 and 2.6'),
('', '13c', 'Describe any methods used to tabulate or visually display results of individual studies and syntheses.', 'Section 2.8'),
('', '13d', 'Describe any methods used to synthesize results and provide a rationale for the choice(s). If meta-analysis was performed, describe the model(s), method(s) to identify the presence and extent of statistical heterogeneity, and software package(s) used.', 'Section 2.8 (structured narrative synthesis; meta-analysis not appropriate, rationale given)'),
('', '13e', 'Describe any methods used to explore possible causes of heterogeneity among study results.', 'Section 2.8 (stratification by validation stage and number of fallers)'),
('', '13f', 'Describe any sensitivity analyses conducted to assess robustness of the synthesized results.', 'Section 2.8'),
('Reporting bias assessment', '14', 'Describe any methods used to assess risk of bias due to missing results in a synthesis (arising from reporting biases).', 'Section 2.8'),
('Certainty assessment', '15', 'Describe any methods used to assess certainty (or confidence) in the body of evidence for an outcome.', 'Section 2.9'),
('RESULTS', None, None),
('Study selection', '16a', 'Describe the results of the search and selection process, from the number of records identified in the search to the number of studies included in the review, ideally using a flow diagram.', 'Section 3.1; Figure 1'),
('', '16b', 'Cite studies that might appear to meet the inclusion criteria, but which were excluded, and explain why they were excluded.', 'Section 3.1; Table S3'),
('Study characteristics', '17', 'Cite each included study and present its characteristics.', 'Section 3.2; Table 2; Table S2'),
('Risk of bias in studies', '18', 'Present assessments of risk of bias for each included study.', 'Section 3.9; Figure 3; Tables S5 and S6'),
('Results of individual studies', '19', 'For all outcomes, present, for each study: (a) summary statistics for each group (where appropriate) and (b) an effect estimate and its precision (e.g., confidence/credible interval), ideally using structured tables or plots.', 'Table 2; Table S2; Figure 2'),
('Results of syntheses', '20a', 'For each synthesis, briefly summarise the characteristics and risk of bias among contributing studies.', 'Sections 3.3–3.9'),
('', '20b', 'Present results of all statistical syntheses conducted. If meta-analysis was done, present for each the summary estimate and its precision (e.g., confidence/credible interval) and measures of statistical heterogeneity. If comparing groups, describe the direction of the effect.', 'Section 3.6; Table 3 (descriptive summaries; no meta-analysis)'),
('', '20c', 'Present results of all investigations of possible causes of heterogeneity among study results.', 'Section 3.6; Table 3; Figure 2'),
('', '20d', 'Present results of all sensitivity analyses conducted to assess the robustness of the synthesized results.', 'Sections 3.3 and 3.9'),
('Reporting biases', '21', 'Present assessments of risk of bias due to missing results (arising from reporting biases) for each synthesis assessed.', 'Section 3.10'),
('Certainty of evidence', '22', 'Present assessments of certainty (or confidence) in the body of evidence for each outcome assessed.', 'Section 3.10; Table 4'),
('DISCUSSION', None, None),
('Discussion', '23a', 'Provide a general interpretation of the results in the context of other evidence.', 'Sections 4.1 and 4.2'),
('', '23b', 'Discuss any limitations of the evidence included in the review.', 'Sections 4.3 and 4.7'),
('', '23c', 'Discuss any limitations of the review processes used.', 'Section 4.7'),
('', '23d', 'Discuss implications of the results for practice, policy, and future research.', 'Sections 4.5 and 4.6; Table 5; Section 5'),
('OTHER INFORMATION', None, None),
('Registration and protocol', '24a', 'Provide registration information for the review, including register name and registration number, or state that the review was not registered.', 'Section 2.1 (PROSPERO CRD420261502037); Abstract'),
('', '24b', 'Indicate where the review protocol can be accessed, or state that a protocol was not prepared.', 'Section 2.1 (PROSPERO record)'),
('', '24c', 'Describe and explain any amendments to information provided at registration or in the protocol.', 'Section 2.10'),
('Support', '25', 'Describe sources of financial or non-financial support for the review, and the role of the funders or sponsors in the review.', 'Funding; Conflicts of Interest'),
('Competing interests', '26', 'Declare any competing interests of review authors.', 'Conflicts of Interest'),
('Availability of data, code and other materials', '27', 'Report which of the following are publicly available and where they can be found: template data collection forms; data extracted from included studies; data used for all analyses; analytic code; any other materials used in the review.', 'Data Availability Statement; Table 2; Tables S1–S6'),
]
d = base_doc()
h = d.add_paragraph(); r = h.add_run('PRISMA 2020 Checklist'); r.bold = True; r.font.size = Pt(14)
p = d.add_paragraph(); r = p.add_run(C.TITLE); r.italic = True; r.font.size = Pt(10)
p = d.add_paragraph(); r = p.add_run('Locations refer to sections, tables, and figures of the main manuscript and to Supplementary Tables S1–S6.'); r.font.size = Pt(9)
rows = [x for x in P]
tb = d.add_table(rows=1 + len(rows), cols=4); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
W = [Cm(3.4), Cm(1.1), Cm(8.4), Cm(3.7)]
for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), W): gc.set(qn('w:w'), str(int(w.twips if hasattr(w,'twips') else w/635)))
hdr = ['Section and topic', 'Item #', 'Checklist item', 'Location where item is reported']
for i, t in enumerate(hdr):
    c = tb.rows[0].cells[i]; c.width = W[i]; c.text = ''; rr = c.paragraphs[0].add_run(t); rr.bold = True; rr.font.size = Pt(8.5); shade(c, 'D9E2F3')
for ri, row in enumerate(rows, start=1):
    cells = tb.rows[ri].cells
    for i in range(4): cells[i].width = W[i]
    if row[1] is None:
        m = cells[0].merge(cells[3]); m.text = ''; rr = m.paragraphs[0].add_run(row[0]); rr.bold = True; rr.font.size = Pt(8.5); shade(m, 'F2F2F2')
        continue
    for i, t in enumerate(row):
        cells[i].text = ''; rr = cells[i].paragraphs[0].add_run(t); rr.font.size = Pt(8.5)
        if i == 0: rr.bold = True
p = d.add_paragraph(); r = p.add_run('From: Page MJ, McKenzie JE, Bossuyt PM, et al. The PRISMA 2020 statement: an updated guideline for reporting systematic reviews. BMJ 2021;372:n71. doi:10.1136/bmj.n71. This work is licensed under CC BY 4.0.'); r.font.size = Pt(8)
d.save('out/PRISMA_2020_Checklist.docx')

# ---------------- Cover letter ----------------
d = base_doc(); d.styles['Normal'].font.size = Pt(10)
def L(text='', bold=False, italic=False, size=None, space_after=6, align=None):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    if align: p.alignment = align
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
    return p
L('Dr. Yosef Hasan Jbara (corresponding author)', bold=True, space_after=0)
L('Computer Engineering Department, College of Engineering & Information Technology, Buraydah Colleges, Buraydah, Saudi Arabia', space_after=0)
L('yosef.hasan@bpc.edu.sa', space_after=12)
L('The Editor-in-Chief and Guest Editor', space_after=0)
L('Diagnostics (MDPI)', italic=True, space_after=0)
L('Special Issue: “Wearable Sensors and Artificial Intelligence for Real-Time Disease Monitoring”', space_after=12)
L('Dear Editors,')
p = L(); p.add_run('We are pleased to submit our manuscript entitled “'); rr = p.add_run(C.TITLE); rr.italic = True; p.add_run('” for consideration as a Systematic Review in ')
rr = p.add_run('Diagnostics'); rr.italic = True; p.add_run(', for the Special Issue “Wearable Sensors and Artificial Intelligence for Real-Time Disease Monitoring”.')
L('Wearable sensors, smartphones, and camera-based systems are increasingly marketed and studied as tools to predict falls in older adults, and many published models report areas under the curve above 0.85. Clinicians, health systems, and device developers need to know whether these prognostic claims are trustworthy. This review answers that question directly.')
L('Why the work matters:', bold=True, space_after=2)
for b in [
 'Scope and rigour. We synthesized 47 reports from 37 independent cohort families, restricting evidence to strictly prospective, individualized fall prediction and grouping overlapping publications from the same cohorts so that the evidence base is not overstated. The protocol was registered in PROSPERO (CRD420261502037), and reporting follows PRISMA 2020 and TRIPOD-SRMA.',
 'First application of PROBAST+AI to this field. We appraised every development and evaluation with the 2025 PROBAST+AI tool, which was designed for both regression and machine-learning models.',
 'A clear, clinically relevant message. Discrimination falls steeply as evaluation becomes more rigorous: the median representative AUC was 0.76, the highest values came from small samples evaluated in their own development data, and the only independent external validation produced AUCs of 0.49–0.59. Calibration was reported by one study, clinical utility by none, and all 47 evaluations were at high risk of bias.',
 'Actionable output. A clinical-readiness scorecard (Table 4) and a set of minimum standards for future studies (Table 5) give developers, reviewers, and regulators concrete criteria for moving digital gait biomarkers from promising signal to validated diagnostic tool.',
]:
    p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3); p.add_run(b)
L('Fit with the Special Issue. The manuscript addresses the clinical validation of AI algorithms applied to wearable-sensor and ambient-vision data for continuous risk monitoring, which is central to the aims of the Special Issue, and it provides a methodological benchmark for future submissions in this area.', space_after=6)
L('Declarations. This manuscript is original, has not been published previously, and is not under consideration by any other journal. All authors have approved the submitted version and agree with its submission to Diagnostics. The review used published data only, so ethics approval was not required. The work was funded by the King Salman Center for Disability Research (KSRG-2026-455); the funder had no role in the study. The authors declare no conflicts of interest. The PRISMA 2020 checklist and Supplementary Tables S1–S6 are provided with the submission.')
L('Thank you for considering our work. We look forward to the reviewers’ comments.', space_after=12)
L('Yours sincerely,', space_after=2)
L('Yosef Hasan Jbara, on behalf of all authors', space_after=0)
d.save('out/Cover_Letter_Diagnostics.docx')
print('ok')
